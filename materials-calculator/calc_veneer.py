# -*- coding: utf-8 -*-
"""
حساب احتياج القشرة
------------------
يقرأ ملف إكسل فيه تابتان:
  - "الاعداد "      : اسم المنتج، العدد المطلوب إنتاجه
  - "احتياج القشرة" : اسم المنتج، ونص حر بالاحتياج مثل «2 متر ارو مسنن»

وينتج ملف إكسل فيه إجمالي الأمتار المطلوبة لكل نوع قشرة.

الاستخدام:
    python calc_veneer.py <ملف_القشرة.xlsx> [ملف_المخرجات.xlsx]
"""

import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SHEET_COUNTS = "الاعداد "
SHEET_NEED = "احتياج القشرة"

# كميات مكتوبة نصًا بدل الأرقام
QTY_WORDS = {"نص": 0.5, "نصف": 0.5, "ربع": 0.25, "واحد": 1.0}

# أسماء في تابة القشرة مكتوبة بصيغة مختلفة عن تابة الأعداد
ALIASES = {
    "palma dresser": "Palma Drawer Dresser",
    "beech side": "Beech Side Tables",
    "mirlia dresser": "Marlia Drawer Dresser",
    "sleek coffebar coffee": "Sleek Coffe Bar",
    "malia tv": "Malia TV unit",
    "fluted round": "Fluted Round Middle Table",
    "rattan corner shelf coffee": "Rattan Corner Shelf",
    "costa rattan desk coffee": "Costa Rattan Desk",
    # سطر واحد يغطي المقاسين معًا — العدد المستخدم هو عدد الأطقم
    "odeno middle table set": "Odeon Middle Table large",
}
SET_ROWS = {"odeno middle table set"}  # أسطر تغطي أكثر من مقاس

UNKNOWN = "غير محدد"


def clean_text(value):
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def match_key(value):
    """مفتاح مطابقة الأسماء: بدون أقواس ولا شرطات سفلية ولا حالة أحرف."""
    text = re.sub(r"[\[\]()_]", " ", clean_text(value))
    return re.sub(r"\s+", " ", text).strip().lower()


def parse_need(text):
    """«2 متر ارو مسنن» -> ("ارو مسنن", 2.0)   |   «نص متر ارو مفجر» -> (..., 0.5)"""
    value = clean_text(text)
    if not value:
        return "", None

    if "مسنن" in value:
        kind = "ارو مسنن"
    elif "مفجر" in value:
        kind = "ارو مفجر"
    elif "امريكي" in value:
        kind = "ارو امريكي"
    else:
        kind = ""

    number = re.match(r"^([\d.]+)", value)
    if number:
        return kind, float(number.group(1))
    for word, amount in QTY_WORDS.items():
        if value.startswith(word):
            return kind, amount
    if value.startswith("متر"):  # «متر ارو مسنن» = متر واحد
        return kind, 1.0
    return kind, None


def load(path):
    need = pd.read_excel(path, sheet_name=SHEET_NEED).iloc[:, :2]
    need.columns = ["product_raw", "need_text"]
    need = need[need["product_raw"].notna()].copy()

    counts = pd.read_excel(path, sheet_name=SHEET_COUNTS).iloc[:, :2]
    counts.columns = ["product", "count"]
    counts = counts[counts["product"].notna()].copy()
    counts["product"] = counts["product"].map(clean_text)
    counts["count"] = pd.to_numeric(counts["count"], errors="coerce").fillna(0)
    return need, counts


def build(need, counts):
    known = {match_key(p): (p, c) for p, c in zip(counts["product"], counts["count"])}

    rows = []
    for _, r in need.iterrows():
        raw = clean_text(r["product_raw"])
        kind, per_unit = parse_need(r["need_text"])
        key = match_key(ALIASES.get(match_key(raw), raw))
        product, count = known.get(key, ("", 0))
        rows.append({
            "product_raw": raw,
            "need_text": clean_text(r["need_text"]),
            "product": product,
            "kind": kind or UNKNOWN,
            "per_unit": per_unit,
            "count": count,
        })

    detail = pd.DataFrame(rows)
    detail["per_unit"] = pd.to_numeric(detail["per_unit"], errors="coerce")
    detail["total_m"] = detail["per_unit"].fillna(0) * detail["count"]
    detail = detail.sort_values(["kind", "total_m"], ascending=[True, False], kind="stable")

    totals = (
        detail.groupby("kind", as_index=False)
        .agg(total_m=("total_m", "sum"),
             products=("product_raw", "count"),
             pieces=("count", "sum"))
        .sort_values("total_m", ascending=False)
    )
    return detail, totals


def build_issues(detail, counts):
    rows = []
    for _, r in detail.iterrows():
        if pd.isna(r["per_unit"]):
            rows.append(["كمية غير مقروءة", r["product_raw"],
                         f"النص «{r['need_text']}» لا يبدأ برقم — لم يُحسب."])
        if r["kind"] == UNKNOWN:
            rows.append(["نوع قشرة غير معروف", r["product_raw"],
                         f"النص «{r['need_text']}» لا يذكر مسنن ولا مفجر."])
        if not r["product"]:
            per = "؟" if pd.isna(r["per_unit"]) else f"{r['per_unit']:g}"
            rows.append(["منتج بلا عدد", r["product_raw"],
                         f"غير موجود في تابة «{SHEET_COUNTS.strip()}» — حُسب بعدد صفر. "
                         f"لو أضفت له عددًا، احتياجه {per} متر {r['kind']} للقطعة."])
        elif match_key(r["product_raw"]) in SET_ROWS:
            rows.append(["سطر يغطي أكثر من مقاس", r["product_raw"],
                         f"سطر واحد للمقاسين. حُسب على أنه {r['count']:g} طقم × "
                         f"{r['per_unit']:g} متر = {r['total_m']:g} متر. لو المقصود "
                         f"{r['per_unit']:g} متر لكل مقاس على حدة، يصبح "
                         f"{r['total_m'] * 2:g} متر."])
        elif match_key(r["product"]) != match_key(r["product_raw"]):
            rows.append(["اسم مكتوب بصيغة مختلفة", r["product_raw"],
                         f"طُوبق مع «{r['product']}» في تابة الأعداد (عدده {r['count']:g})."])

    listed = {match_key(p) for p in detail["product"] if p}
    for product, count in zip(counts["product"], counts["count"]):
        if match_key(product) not in listed:
            rows.append(["منتج بلا سطر قشرة", product,
                         f"مطلوب {count:g} قطعة ولا يوجد له سطر في تابة "
                         f"«{SHEET_NEED}» — اعتُبر بلا قشرة."])
    return pd.DataFrame(rows, columns=["النوع", "البند", "التفاصيل"])


# --------------------------------------------------------------------------
# كتابة الملف
# --------------------------------------------------------------------------
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
IN_HEAD_FILL = PatternFill("solid", fgColor="7F6000")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)


def write_sheet(wb, title, headers, rows, widths, number_cols=(),
                table_name=None, head_fill=HEAD_FILL):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = head_fill
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for row in rows:
        ws.append(row)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.column in number_cols:
                cell.number_format = "#,##0.00"

    if ws.max_row > 1:
        ws.freeze_panes = "A2"
        if table_name:
            table = Table(displayName=table_name,
                          ref=f"A1:{get_column_letter(len(headers))}{ws.max_row}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        else:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    return ws


def write_output(path, detail, totals, issues, counts):
    wb = Workbook()
    wb.remove(wb.active)

    grand = float(totals["total_m"].sum())
    ws = write_sheet(
        wb, "الإجمالي",
        ["نوع القشرة", "إجمالي الأمتار المطلوبة", "عدد المنتجات", "إجمالي القطع"],
        [[r["kind"], round(float(r["total_m"]), 2), int(r["products"]), int(r["pieces"])]
         for _, r in totals.iterrows()] +
        [["الإجمالي العام", round(grand, 2), int(totals["products"].sum()), ""]],
        [24, 26, 15, 15], number_cols={2},
    )
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")

    write_sheet(
        wb, "تفصيل المنتجات",
        ["المنتج", "نوع القشرة", "متر للوحدة", "العدد المطلوب",
         "إجمالي الأمتار", "النص الأصلي"],
        [[r["product"] or r["product_raw"], r["kind"],
          "—" if pd.isna(r["per_unit"]) else float(r["per_unit"]),
          int(r["count"]), round(float(r["total_m"]), 2), r["need_text"]]
         for _, r in detail.iterrows()],
        [32, 16, 14, 15, 16, 24], number_cols={3, 5}, table_name="VeneerDetail",
    )

    write_sheet(
        wb, "ملاحظات", ["النوع", "البند", "التفاصيل"],
        issues.values.tolist(), [26, 32, 95], table_name="VeneerIssues",
    )
    for row in wb["ملاحظات"].iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    write_sheet(
        wb, "مدخلات - احتياج القشرة",
        ["المنتج (كما ورد)", "النص الأصلي", "نوع القشرة", "متر للوحدة"],
        [[r["product_raw"], r["need_text"], r["kind"],
          "" if pd.isna(r["per_unit"]) else float(r["per_unit"])]
         for _, r in detail.iterrows()],
        [32, 24, 16, 14], number_cols={4},
        table_name="InputNeed", head_fill=IN_HEAD_FILL,
    )
    write_sheet(
        wb, "مدخلات - الأعداد", ["اسم المنتج", "العدد"],
        [[p, int(c)] for p, c in zip(counts["product"], counts["count"])],
        [40, 14], number_cols={2}, table_name="InputCounts", head_fill=IN_HEAD_FILL,
    )
    wb.save(path)


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else "احتياج_القشرة.xlsx"

    need, counts = load(src)
    detail, totals = build(need, counts)
    issues = build_issues(detail, counts)
    write_output(dst, detail, totals, issues, counts)

    print(f"تم إنشاء: {dst}")
    for _, row in totals.iterrows():
        print(f"  {row['kind']:<12}: {row['total_m']:>10,.2f} متر  "
              f"({int(row['products'])} منتج)")
    print(f"  {'الإجمالي':<12}: {totals['total_m'].sum():>10,.2f} متر")
    print(f"  ملاحظات للمراجعة: {len(issues)}")


if __name__ == "__main__":
    main()
