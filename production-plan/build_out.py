# -*- coding: utf-8 -*-
import runpy, sys, math, datetime as dt, openpyxl
from collections import defaultdict, Counter
from openpyxl.styles import Font, PatternFill, Alignment
g=runpy.run_path('sched.py')
ORDERS=g['ORDERS']; DAYS=g['DAYS']; W=g['WORKERS']; NORM=g['NORM_H']; OT=g['OT_H']
run=g['run']; ops=g['ops']; name2code=g['name2code']; nz=g['nz']

def finish_map(done,hist):
    fin={}
    for o in ORDERS:
        k=o['rank']; last=-1; ok=True
        for d,Wd in o['work'].items():
            if done[k][d] < Wd-1e-6: ok=False;break
            last=max(last,max(t for t,_ in hist[k][d]))
        fin[k]=DAYS[last] if ok and last>=0 else None
    return fin
def dept_span(hist,k):
    out={}
    for d,l in hist[k].items():
        if l: out[d]=(DAYS[min(t for t,_ in l)], DAYS[max(t for t,_ in l)], sum(h for _,h in l))
    return out

R={}
R['واقعي']=run(False);            F_real=finish_map(R['واقعي'][0],R['واقعي'][1])
R['متفائل']=run(True);            F_opt =finish_map(R['متفائل'][0],R['متفائل'][1])
for o in ORDERS: o['_w']=o['work']; o['work']={d:h*1.15 for d,h in o['work'].items()}
R['متشائم']=run(False);           F_pes =finish_map(R['متشائم'][0],R['متشائم'][1])
for o in ORDERS: o['work']=o['_w']
done,hist,deptday,plan=R['واقعي']
doneO,histO,deptdayO,planO=R['متفائل']
print("واقعي: P1 آخر", max(F_real[o['rank']] for o in ORDERS if o['P']==1),
      "| P2 آخر", max(F_real[o['rank']] for o in ORDERS if o['P']==2))
print("متشائم: P1 آخر", max(F_pes[o['rank']] for o in ORDERS if o['P']==1),
      "| P2 آخر", max(F_pes[o['rank']] for o in ORDERS if o['P']==2))

# ===== طابور CNC =====
CNCJOB=defaultdict(float)   # (يوم, منتج) -> دقايق ماكينة
cnc_by_order={}
for o in ORDERS:
    code=name2code.get(o['prod'])
    if not code: continue
    m=0
    for dp in ('نجارة التنجيد','نجارة النوم والسفرة'):
        for sq,nm,mn in ops.get(code,{}).get(dp,[]):
            if 'cnc' in nz(nm) or 'سي ان' in nz(nm): m+=mn
    if m>0 and dp in o['work']: cnc_by_order[o['rank']]=min(max(m,20),90)   # ⅓ س ← 1.5 س
for o in ORDERS:
    k=o['rank']
    if k not in cnc_by_order: continue
    for d in ('نجارة التنجيد','نجارة النوم والسفرة'):
        if d in hist[k] and hist[k][d]:
            t=min(t for t,_ in hist[k][d]); CNCJOB[(t,o['prod'])]+=cnc_by_order[k]; break
cnc_rows=[]; CNC_DAY=8*60
byday=defaultdict(list)
for (t,p),m in CNCJOB.items(): byday[t].append((p,m))
carry=0
for t in sorted(byday):
    jobs=byday[t]; tot=sum(m for _,m in jobs)
    # تجميع نفس المنتج = وفر 30% من زمن التجهيز
    grouped=sum(m*0.7 if len([1 for pp,_ in jobs if pp==p])>1 else m for p,m in jobs)
    q=max(0,carry+grouped-CNC_DAY); use=min(CNC_DAY,carry+grouped)
    cnc_rows.append([str(DAYS[t]),len(jobs),round(tot/60,2),round(grouped/60,2),round(use/60,2),
                     f"{100*use/CNC_DAY:.0f}%",round(q/60,2),'، '.join(sorted({p for p,_ in jobs})[:4])])
    carry=q
print(f"CNC: أيام فيها شغل {len(cnc_rows)} | أقصى انتظار {max((r[6] for r in cnc_rows),default=0):.1f} س")
g2=globals()

# ================== بناء الوركبوك ==================
wb=openpyxl.Workbook(); wb.remove(wb.active)
HF=Font(bold=True,color='FFFFFF'); HB=PatternFill('solid',fgColor='31538F')
RED=PatternFill('solid',fgColor='FFCDD2'); YEL=PatternFill('solid',fgColor='FFF9C4')
ORG=PatternFill('solid',fgColor='FFE0B2'); GRN=PatternFill('solid',fgColor='C8E6C9')
def sh(n,h,rows,w):
    ws=wb.create_sheet(n); ws.sheet_view.rightToLeft=True; ws.append(h)
    for c in ws[1]: c.font=HF;c.fill=HB;c.alignment=Alignment(horizontal='center',wrap_text=True)
    for r in rows: ws.append(r)
    for i,x in enumerate(w,1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=x
    ws.freeze_panes='A2'; return ws

# 1 SETTINGS
sh('SETTINGS',['البند','القيمة','ملاحظة'],[
 ['تاريخ بداية الخطة','2026-09-06',''],['تاريخ توفر الخامات','2026-09-07','بكرة'],
 ['الوردية الصباحية','7:30 ← 12:30 = 5 س',''],['البريك','12:30 ← 1:30','مش إنتاج'],
 ['بعد البريك','1:30 ← 4:30 = 3 س',''],['اليوم العادي','8 ساعات',''],
 ['السهر','4:30 ← 9:00 = 4.5 س','حد أقصى للفرد'],['أقصى يوم كامل','12.5 ساعة',''],
 ['الجمعة','إجازة','متغيّر — غيّرها لو في ضغط'],
 ['كفاءة كل قسم','75%','تخمين معتمد منك'],['هامش أمان','10%','أعطال/غياب/إعادة شغل'],
 ['معامل الصباح','100%','منحنى اليوم'],['معامل بعد البريك','90%','منحنى اليوم'],
 ['معامل السهر','70%','إرهاق'],
 ['طاقة الفرد/اليوم عادي',round(NORM,2),'(5×1.00 + 3×0.90) × 75% × 90%'],
 ['طاقة الفرد/اليوم سهر',round(OT,2),'4.5 × 0.70 × 75% × 90%'],
 ['خصم التكرار 5-9 وحدة','10%',''],['خصم التكرار 10-19','15%',''],['خصم التكرار 20+','20%',''],
 ['تجفيف السيلر','ليلة كاملة','الشغل يكمّل صباح اليوم التالي'],
 ['تجفيف الدوكو','2-3 ساعات','يخرج نفس اليوم لو دخل بدري'],
 ['ماكينات CNC','1','مورد مشترك — طابور مستقل'],
 ['زمن CNC للشغلانة','⅓ ← 1.5 ساعة',''],
 ['سيناريو متفائل','بالسهر',''],['سيناريو واقعي','ساعات عادية بس','ده اللي أعتمد عليه'],
 ['سيناريو متشائم','عادي + 15% وقت','إعادة شغل/غياب'],
 ['سقف الطاقة','مايتعداش','الزيادة تترحّل — بأمرك'],
 ['مستبعد','Kosan / التغليف / تانجل / هاربر ترابيزة / لوليتا جانبية','بأمرك'],
],[30,42,44])

# 2 ORDERS
orow=[[o['id'],o['prod'],o['qty'],o['trk'][:44],f"P{o['P']}",o['lab'],
       str(o['due']) if o['due'] else '—',o['cur'] or '—',str(o['stage'] or '—'),
       o['src'],round(sum(o['work'].values()),2)] for o in ORDERS]
w=sh('ORDERS',['كود الأمر','المنتج','الكمية','العميل / كود التتبّع','الأولوية','المجموعة',
     'التسليم المطلوب','القسم الحالي','المرحلة الحالية','مصدر الوقت','ساعات فاضلة'],orow,[18,30,7,42,8,30,14,20,24,26,12])
for r in w.iter_rows(min_row=2):
    if r[4].value=='P1':
        for c in r: c.fill=YEL

# 3 SCHEDULE
srow=[]
for o in ORDERS:
    for d,(s,e,h) in sorted(dept_span(hist,o['rank']).items(), key=lambda x:x[1][0]):
        srow.append([o['id'],o['prod'],f"P{o['P']}",d,str(s),str(e),round(h,2),(e-s).days+1])
sh('SCHEDULE',['كود الأمر','المنتج','الأولوية','القسم','تاريخ البداية','تاريخ النهاية','ساعات','مدى الأيام'],srow,[18,30,8,22,14,14,9,10])

# 4 DAILY_PLAN
drow=[]
for t,day in enumerate(DAYS):
    if not deptday[t]: continue
    if day> max(F_real.values()): break
    for d,h in sorted(deptday[t].items(), key=lambda x:-x[1]):
        items=sorted(plan[t][d], key=lambda x:-x[1])[:6]
        nm='، '.join(f"{ORDERS[k-1]['prod'][:22]} ({hh:.1f}س)" for k,hh in items)
        need=math.ceil(h/NORM); otp=max(0,need-W[d])
        drow.append([str(day),['الاتنين','الثلاثاء','الأربعاء','الخميس','الجمعة','السبت','الأحد'][day.weekday()],
                     d,round(h,1),W[d],round(W[d]*NORM,1),f"{100*h/(W[d]*NORM):.0f}%",
                     len(plan[t][d]),nm])
w=sh('DAILY_PLAN',['التاريخ','اليوم','القسم','ساعات مطلوبة','عدد العمال','الطاقة العادية','التحميل %','عدد الأوردرات','أهم الأوردرات في اليوم'],drow,[13,10,22,13,11,14,11,13,72])
for r in w.iter_rows(min_row=2):
    v=float(str(r[6].value).replace('%',''))
    if v>=99:
        for c in r: c.fill=ORG

# 5 DEPT_LOAD
lrow=[]
for t,day in enumerate(DAYS):
    if day>max(F_real.values()): break
    for d in sorted(W):
        h=deptday[t][d]; cap=W[d]*NORM
        lrow.append([str(day),d,round(h,1),round(cap,1),f"{100*h/cap:.0f}%",
                     round(max(0,h-cap),1),'مختنق' if h>=cap*.99 else ('فاضي' if h<cap*.3 else 'عادي')])
w=sh('DEPT_LOAD',['التاريخ','القسم','ساعات مطلوبة','الطاقة المتاحة','التحميل %','عجز','الحالة'],lrow,[13,22,13,14,11,9,11])
for r in w.iter_rows(min_row=2):
    if r[6].value=='مختنق':
        for c in r: c.fill=RED
    elif r[6].value=='فاضي':
        for c in r: c.fill=GRN

# 6 CAPACITY
crow=[]
for d in sorted(W,key=lambda x:-sum(deptday[t][x] for t in range(len(DAYS)))):
    tot=sum(deptday[t][d] for t in range(len(DAYS)))
    busy=sum(1 for t in range(len(DAYS)) if deptday[t][d]>=W[d]*NORM*0.99)
    crow.append([d,W[d],8,'75%','10%',round(NORM,2),round(W[d]*NORM,1),round(W[d]*OT,1),
                 round(tot,0),round(tot/(W[d]*NORM),1),busy])
sh('CAPACITY',['القسم','عدد العمال','ساعات خام/فرد','الكفاءة','هامش الأمان','طاقة الفرد الفعلية',
   'طاقة القسم/يوم','طاقة السهر/يوم','إجمالي الشغل (ساعة)','أيام شغل','أيام محمّلة 100%'],crow,[22,11,13,10,12,16,15,15,17,10,16])

# 7 CNC_QUEUE
sh('CNC_QUEUE',['التاريخ','عدد الشغلانات','زمن خام (س)','بعد تجميع المتشابه','المستخدم من الماكينة','التحميل %','منتظر لبكرة (س)','المنتجات'],cnc_rows,[13,13,13,18,18,11,15,50])

# 8 OVERTIME
otrow=[]
for t,day in enumerate(DAYS):
    if day>max(F_opt.values()): break
    for d in sorted(W):
        used=deptdayO[t][d]; capn=W[d]*NORM
        if used>capn+0.01:
            exc=used-capn; ppl=min(W[d],math.ceil(exc/OT))
            reason=sorted(planO[t][d],key=lambda x:-x[1])[:2]
            otrow.append([str(day),d,ppl,W[d],round(exc,1),round(ppl*4.5,1),
                          '، '.join(ORDERS[k-1]['prod'][:26] for k,_ in reason)])
w=sh('OVERTIME',['التاريخ','القسم','عدد اللي يسهروا','إجمالي عمال القسم','العجز (ساعة فعلية)','ساعات السهر الحقيقية','سبب السهر (أنهي أوردر)'],otrow,[13,22,15,17,18,19,52])
for r in w.iter_rows(min_row=2):
    if r[2].value==r[3].value:
        for c in r: c.fill=RED

# 9 DELIVERY
vrow=[]
for o in ORDERS:
    k=o['rank']; a,b,c_=F_opt[k],F_real[k],F_pes[k]
    late=(b-o['due']).days if o['due'] and b else None
    span=(c_-a).days if a and c_ else 0
    conf='عالية' if o['src']=='مسار من الملف' and span<=3 else ('متوسطة' if span<=7 else 'منخفضة')
    if o['src'].startswith('تقديري'): conf='منخفضة'
    vrow.append([o['id'],o['prod'],o['qty'],f"P{o['P']}",o['lab'],str(o['due']) if o['due'] else '—',
                 str(a),str(b),str(c_),late if late is not None else '—',conf,o['src']])
vrow.sort(key=lambda r:(r[3],r[7]))
w=sh('DELIVERY',['كود الأمر','المنتج','الكمية','الأولوية','المجموعة','التسليم المطلوب',
     'متفائل','واقعي (اعتمد عليه)','متشائم','متأخر (يوم)','الثقة','مصدر الوقت'],vrow,[18,30,7,8,28,14,12,17,12,12,10,26])
for r in w.iter_rows(min_row=2):
    if isinstance(r[9].value,int) and r[9].value>0:
        for c in r: c.fill=RED
    elif r[10].value=='منخفضة':
        for c in r: c.fill=ORG

# 10 FLOW
frow=[]
for d in sorted(W):
    ts=[t for t in range(len(DAYS)) if deptday[t][d]>0.01]
    if not ts: continue
    gaps=[t for t in range(min(ts),max(ts)+1) if deptday[t][d]<W[d]*NORM*0.3]
    full=[t for t in range(min(ts),max(ts)+1) if deptday[t][d]>=W[d]*NORM*0.99]
    frow.append([d,str(DAYS[min(ts)]),str(DAYS[max(ts)]),len(ts),len(full),len(gaps),
                 f"{100*len(full)/max(1,len(ts)):.0f}%",
                 'اختناق — شغّال 100% طول الوقت' if len(full)/max(1,len(ts))>.8 else
                 ('فيه فجوات — ينفع ننقل منه عمال' if len(gaps)>len(ts)*.3 else 'متوازن'),
                 '، '.join(str(DAYS[t]) for t in gaps[:6])])
sh('FLOW',['القسم','أول يوم شغل','آخر يوم شغل','أيام شغل','أيام محمّل 100%','أيام شبه فاضي','نسبة التحميل الكامل','التقييم','أمثلة أيام فاضية'],frow,[22,14,14,11,16,14,17,34,52])

# 11 ISSUES
sh('ISSUES',['المشكلة','التفاصيل','الأثر على الخطة'],[
 ['ملف المسارات الجديد مطابق للقديم 100%','بصمة المحتوى واحدة — مفيش تعديل اتعمل','لسه محتاج النسخة المعدّلة'],
 ['9 صفوف من غير وقت معياري','روتري سرير (8 عمليات) + تريكس فوتيه ("لا وقت")','الأوردرات دي وقتها أقل من الحقيقة'],
 ['319 وحدة بوقت تقديري','متوسط منتجات مشابهة في نفس الفئة','ثقة منخفضة في شيت DELIVERY'],
 ['8 منتجات فيها استانلس + قشرة + سفنجة','هالاند ثنائية/ثلاثية، زين كرسي، اركي سرير، نولا بوفيه، روتري تسريحة، نيفادا كومود، لوتي ترابيزة','اعتبرتهم 3 فروع متوازية — محتاج تأكيد'],
 ['عدد كابينات الدهان والمجففات','مش معروف','التجفيف محسوب كوقت تقويم بس، من غير طابور كابينة'],
 ['تواريخ P2','562 أمر مالهاش تاريخ تسليم','رتّبتهم بترتيب ملف الأولويات زي ما قلتي'],
 ['⚠ ماتشا فوتيه — القماش لسه متحددش','7 وحدات مجدولة على الكسوة يوم 7/9','لو القماش ماتحددش، الكسوة هتقف والتاريخ هيتأخر'],
 ['⚠ بابلي فوتيه — 7 وحدات قماشهم يتحدد بكرة','اتجدولوا من 8/9 مش 7/9','التسعة التانية شغّالين عادي من 7/9'],
 ['أوامر BPO مش هتشتغل','194 وحدة اتشالت (كل الأكواد اللي بتبدأ BPO)','ماعدا BPO-000055 = الـ128 كرسي سبايدر بتوع عمرو حسن — دول فاضلين'],
 ['كراسي جنا حازم اتشالت','9 كراسي زين سفرة (8 على NF-003308 + 1 على FXN-000499) — اتشالوا بكود الأمر','بأمرك. كراسي زين بتاعة محمد وحيد (6) لسه في الخطة'],
 ['نيو سيول كرسي سفرة مؤجّلة','22 وحدة — مش هتشتغل الأسبوع ده، أول يوم 12/9','بأمرك'],
 ['اوردين ترابيزة كبيرة خلصت','5 وحدات اتشالت من الجدولة','بأمرك'],
 ['قاعدة: الترابيزات والمرايا والكومودات والتسريحات والبوفيهات والسفرة = نجارة نوم وسفرة','ماعدا اتووم كومود. وكل شغل "عاصم" كمان نوم وسفرة','نقل شغل النجارة من التنجيد للنوم والسفرة في التقديرات'],
 ['تورين فوتيه — وقت تقديري','مالهاش مسار؛ استخدمت متوسط فوتيه (اريكة) والدهانات بعد خصم "خدمة"','ثقة منخفضة'],
],[38,90,44])
wb.save('PRODUCTION_PLAN.xlsx')
print("\nتم بناء PRODUCTION_PLAN.xlsx")
print(f"  SCHEDULE {len(srow)} صف | DAILY_PLAN {len(drow)} | DEPT_LOAD {len(lrow)} | OVERTIME {len(otrow)} | DELIVERY {len(vrow)}")
