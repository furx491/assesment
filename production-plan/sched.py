# -*- coding: utf-8 -*-
"""جدولة إنتاج بطاقة محدودة — تدفق قطعة قطعة + تجميع الفروع + طابور CNC"""
import openpyxl, re, json, math, datetime as dt
from collections import defaultdict, Counter
from override import OVERRIDE, ov_minutes
from status_update import STATUS, STATUS_BY_ID, RT, FINISHED, DEFERRED, EXCLUDE_IDS, MISSING_TIME, is_bpo_excluded, remap, to_nw
U='/root/.claude/uploads/fdb1b8fa-74fd-56bd-9c54-782073111ccf/'
def num(v): return v if isinstance(v,(int,float)) else 0
def nz(s):
    s=str(s).strip().lower(); s=re.sub('[إأآا]','ا',s); s=re.sub('[ىي]','ي',s)
    s=re.sub('ة','ه',s); s=re.sub(r'\(.*?\)',' ',s); s=re.sub(r'[^\w\s]',' ',s)
    return re.sub(r'\s+',' ',s).strip()

# ============ SETTINGS ============
START=dt.date(2026,9,7); EFF=.75; BUF=.10   # الاتنين
NORM_H=(5*1.00+3*0.90)*EFF*(1-BUF)      # 5.20 ساعة/فرد/يوم
OT_H  = 4.5*0.70      *EFF*(1-BUF)      # 2.13 ساعة/فرد/يوم
FRIDAY_OFF=False   # الجمعة شغل — بأمر المستخدمة
SEALER_LAG=1          # رش السيلر آخر اليوم → أقرب عملية تالية صباح اليوم التالي
MAT_DATE=dt.date(2026,9,7)
CNC_MIN_PER_JOB=20; CNC_MAX_MIN=90
# نجارة السراير كرو مستقل: 1 تقديم + 3 تجميع = 4
# مخصومين من نجارة التنجيد (42 ← 38) عشان مانخترعش طاقة زيادة
WORKERS={'نجارة التنجيد':38,'نجارة سراير':4,'السفنجة':19,'التفصيل والكسوة':22,'الاستانلس':16,
         'الدهانات':30,'القشرة':9,'نجارة النوم والسفرة':15,'تشطيب التنجيد':5,'تشطيب نوم وسفرة':3}
BED_DEPT='نجارة سراير'
def is_bed(prod): return ('سرير' in str(prod)) or ('bed' in str(prod).lower())
def bed_route(work, prod):
    """شغل نجارة السراير يروح لقسم السراير المستقل"""
    if not is_bed(prod): return work
    w={}
    for d,h in work.items():
        w2 = BED_DEPT if d in ('نجارة التنجيد','نجارة النوم والسفرة') else d
        w[w2]=w.get(w2,0)+h
    return w
# طاقات فرعية جوه التفصيل والكسوة (البرومبت: خياطة 2، تفصيل 6، فايبر وكسوة 14)
SUBCAP={'التفصيل والكسوة':{'مبكر':(2+6),'متأخر':14}}
EXCLUDE={'Kosan','تانجل ترابيزه جانبيه','هاربر ترابيزة رئيسية اريكة','لوليتا ترابيزة جانبية'}
EXCLUDE|=FINISHED   # خلصت خالص — بأمر المستخدمة

def workdays(n0):
    d=START; out=[]
    while len(out)<n0:
        if not (FRIDAY_OFF and d.weekday()==4): out.append(d)
        d+=dt.timedelta(days=1)
    return out
DAYS=workdays(140); DIDX={d:i for i,d in enumerate(DAYS)}

# ============ المسارات ============
wr=openpyxl.load_workbook(U+'7b215713-_________________....xlsx',data_only=True)['Sheet1']
ops=defaultdict(lambda: defaultdict(list)); pname={}
for r in [x for x in wr.iter_rows(values_only=True)][1:]:
    if not r[3]: continue
    pname[r[3]]=str(r[5]).strip()
    if r[1]=='التغليف': continue
    ops[r[3]][r[1]].append((r[6] if isinstance(r[6],int) else 999,str(r[7] or ''),num(r[8])))
for c in ops:
    for d in ops[c]: ops[c][d].sort()
# الأوقات الناقصة: وزّع الإجمالي المعتمد بالتساوي على عمليات القسم
for c in list(ops):
    for d in list(ops[c]):
        key=(pname.get(c),d)
        if key in MISSING_TIME and sum(x[2] for x in ops[c][d])==0:
            tot,_src=MISSING_TIME[key]; n=len(ops[c][d])
            ops[c][d]=[(sq,nm,tot/n) for sq,nm,_ in ops[c][d]]
name2code={}
for c,n in pname.items(): name2code.setdefault(n,c)
wc=openpyxl.load_workbook(U+'485a723f-______________________.xlsx',data_only=True)['Sheet1']
DM={'استالس':'الاستانلس','نجارة تنجيد':'نجارة التنجيد','سفنجة':'السفنجة','كسوة':'التفصيل والكسوة'}
chair=defaultdict(list)
for r in list(wc.iter_rows(values_only=True))[2:]:
    if r[0] and r[1] in DM: chair[DM[r[1]]].append((r[0].strip(),num(r[2])))

ORDW={'نجارة التنجيد':1,'نجارة النوم والسفرة':1,'الاستانلس':1,'السفنجة':2,'القشرة':2,
      'التفصيل والكسوة':3,'الدهانات':3,'تشطيب التنجيد':4,'تشطيب نوم وسفرة':4}
def branches(code):
    d=set(ops[code]); br={}; split=('السفنجة' in d and 'القشرة' in d)
    for dp in d:
        if dp in ('تشطيب التنجيد','تشطيب نوم وسفرة'): br[dp]=('Z',9)
        elif dp=='الاستانلس' and len(d)>1: br[dp]=('C',1)
        elif dp==BED_DEPT: br[dp]=('A',1)
        elif split and dp in ('القشرة','الدهانات'): br[dp]=('B',ORDW[dp])
        else: br[dp]=('A',ORDW.get(dp,9))
    return br
def remaining(code,cur,stage,fs,oplist=None):
    out=defaultdict(float); br=branches(code)
    if oplist is None: oplist={}
    if fs or not cur or cur not in ops[code]:
        for dp,l in ops[code].items():
            out[dp]=sum(x[2] for x in l); oplist[dp]=[(n,m) for _,n,m in l]
        return out
    cb,cp=br[cur]; l=ops[code][cur]; i=0
    if stage:
        import difflib; s=nz(stage); bs=0; bi=0
        for j,(sq,nm,mn) in enumerate(l):
            n2=nz(nm); rr=1.0 if (s and (s in n2 or n2 in s)) else difflib.SequenceMatcher(None,s,n2).ratio()
            if rr>bs: bs,bi=rr,j
        if bs>=.60: i=bi
    out[cur]=sum(x[2] for x in l[i:]); oplist[cur]=[(n,m) for _,n,m in l[i:]]
    for dp,ld in ops[code].items():
        if dp==cur: continue
        b,p=br[dp]
        if b=='Z' or b!=cb or p>cp:
            out[dp]=sum(x[2] for x in ld); oplist[dp]=[(n,m) for _,n,m in ld]
    return out

# ============ الأوردرات + الأولويات ============
DEPTNAME={'نجارة التنجيد','نجارة النوم والسفرة','التفصيل والكسوة','الدهانات','القشرة','السفنجة',
          'الاستانلس','الفايبر جلاس','دهانات','نجارة','التخطيط'}
CATS=[('كرسي سفرة',['كرسي سفره','كراسي','كرسي']),('نيو بيكا كنبة',['بيكا كنبه مقاس','بيكا كنبه']),
      ('كنبة ثلاثية',['كنبه ثلاثيه','ثلاثيه']),('كنبة ثنائية',['كنبه ثنائيه','ثنائيه','double sofa']),
      ('ركنة',['ركنه','حرف l','حرف u']),('فوتيه',['فوتيه','فوتية','ارمشير']),
      ('بف',['بف','بانكت','بانكيت','رول']),('سرير',['سرير','bed']),('كومود',['كومود','commode']),
      ('تسريحة',['تسريحه','درسينج','دريسنج','درسيينج','نيش','جزامه','فاترينه']),
      ('بوفيه',['بوفيه','يوفيه']),('ترابيزة سفرة',['ترابيزه سفره','سفره']),
      ('ترابيزة جانبية',['ترابيزه جانبيه','سايد','ترابيزه صغيره','side','جانبيه']),
      ('ترابيزة رئيسية',['ترابيزه رئيسيه','ترابيزه كبيره','ديننج','dinning','ترابيزه','coffe','table']),
      ('كونصول',['كونصول','تي في','تى فى','تي فيونت','ديسك','desk']),
      ('مرايا',['مرايا','مرايه','برواز','mirror']),('بلوك',['بلوك']),('كنبة',['كنبه','sofa'])]
def cat_of(n):
    x=nz(n)
    for c,ks in CATS:
        for k in ks:
            if k in x: return c
catrows=defaultdict(list)
for c in ops:
    k=cat_of(pname[c])
    if k: catrows[k].append({d:sum(x[2] for x in l) for d,l in ops[c].items()})
catavg={};catn={}
for k,l in catrows.items():
    agg=defaultdict(float); pres=Counter()
    for d in l:
        for dp,m in d.items(): agg[dp]+=m; pres[dp]+=1
    catavg[k]={dp:agg[dp]/pres[dp] for dp in agg if pres[dp]>=max(1,len(l)*.5)}; catn[k]=len(l)

wo=openpyxl.load_workbook(U+'70dac0e1-_____________________________.xlsx',data_only=True)['أوامر التصنيع']
raw=[r for r in list(wo.iter_rows(values_only=True))[1:] if any(v is not None for v in r)]
active=[r for r in raw if r[4] in ('IN_PROGRESS','NEW','PLANNED')]
wp=openpyxl.load_workbook(U+'738ab119-___________________.xlsx',data_only=True)['خطة الانتاج ']
DAYNM={'السبت':dt.date(2026,9,5),'الاحد':dt.date(2026,9,6),'الأحد':dt.date(2026,9,6),
       'الاتنين':dt.date(2026,9,7),'الثلاثاء':dt.date(2026,9,8),'الأربعاء':dt.date(2026,9,9),
       'الاربعاء':dt.date(2026,9,9),'الخميس':dt.date(2026,9,10)}
due={};cust={};day=None
for r in wp.iter_rows(min_row=5,values_only=True):
    d=str(r[1] or '').split('\n')[0].strip()
    if d in DAYNM: day=DAYNM[d]
    for m in re.findall(r'NF-?(\d{3,6})',str(r[2] or '')):
        c='NF-'+m.zfill(6); due.setdefault(c,day); cust.setdefault(c,str(r[3] or '').strip())
def nfc(t):
    m=re.search(r'[Nn][Ff]-?(\d{3,6})',str(t or '')); return 'NF-'+m.group(1).zfill(6) if m else None
qty=Counter()
for r in active:
    p=str(r[1] or '').strip()
    if p not in EXCLUDE and r[0] not in EXCLUDE_IDS and not is_bpo_excluded(r[0]): qty[p]+=r[2] or 1
def bdisc(q): return .20 if q>=20 else .15 if q>=10 else .10 if q>=5 else 0.0
# متوسط دقائق القسم لكل وحدة عبر كل المنتجات اللي بتعدي عليه — شبكة أمان
DEPT_AVG={}
for _d in set(dd for c in ops for dd in ops[c]):
    _v=[sum(x[2] for x in ops[c][_d]) for c in ops if _d in ops[c]]
    DEPT_AVG[_d]=sum(_v)/len(_v) if _v else 0.0

ORDERS=[]; _seen_status=set()
for r in active:
    p=str(r[1] or '').strip()
    if p in EXCLUDE or r[0] in EXCLUDE_IDS or is_bpo_excluded(r[0]): continue
    q=r[2] or 1; trk=str(r[3] or ''); c=nfc(trk); dsc=bdisc(qty[p])
    stage=r[5]; dept=r[6]; _note=''
    if r[0] in STATUS_BY_ID:
        _sb=STATUS_BY_ID[r[0]]; dept=_sb['cur']; stage=_sb['stage']; _note=_sb.get('note','')
    fs=(r[4]=='NEW') or (stage is None) or (str(stage).strip() in DEPTNAME)
    src='مسار من الملف'; OPL={}
    _skeys=[kk for kk,vv in STATUS.items() if (vv.get('match') or kk)==p]
    if _skeys:
        if p in _seen_status: continue      # الكميات في التحديث إجمالية — أوامر مجمّعة
        _seen_status.add(p)
        for _sk in _skeys:
            st=STATUS[_sk]; q=st['qty']; dsc=bdisc(q)
            _rem=st['rem']
            if _rem is None:                                  # ارجع لتقدير الفئة
                kc=cat_of(p); _cur=ORDW.get(st['cur'],0)
                _rem={d:m for d,m in catavg.get(kc,{}).items() if ORDW.get(d,9)>=_cur}
            work={d:m*q*(1-dsc) for d,m in _rem.items()}
            work=remap(work,p,cat_of(p))
            br={}
            for d in work:
                if d in ('تشطيب التنجيد','تشطيب نوم وسفرة'): br[d]=('Z',9)
                elif d in ('القشرة','الدهانات'): br[d]=('B',ORDW[d])
                else: br[d]=('A',ORDW.get(d,9))
            OPL2={}
            for d in work:
                if p in RT and d in RT[p] and st['rem']:
                    tot=sum(x[2] for x in RT[p][d]) or 1
                    OPL2[d]=[(n,m*q*(1-dsc)*(st['rem'][d]/tot)/60.0) for _,n,m in RT[p][d]]
            if 'اريكة' in p or 'اريكة' in trk: P,sub,lab,dd=1,0,'أريكة — تجاري عاجل',dt.date(2026,9,10)
            elif c and c in due:                P,sub,lab,dd=1,1,f'معرض — {cust.get(c,"")}',due[c]
            elif r[7]=='الاسبوع الحالي':        P,sub,lab,dd=1,2,'مصنع — الأسبوع الحالي',dt.date(2026,9,10)
            else:                               P,sub,lab,dd=2,3,'مصنع — باقي التشغيل',None
            ORDERS.append(dict(id=r[0],prod=_sk,qty=q,trk=trk,nf=c,P=P,sub=sub,lab=lab,due=dd,opl=OPL2,
                work={k2:v2/60.0 for k2,v2 in work.items()},br=br,src='تحديث حالة منك',
                cur=st['cur'],stage=st['stage'],ready=st.get('ready',START)))
        continue
    if False:
        st=None; q=r[2] or 1; dsc=bdisc(qty[p]); work={}
        br={}
        for d in work:
            if d in ('تشطيب التنجيد','تشطيب نوم وسفرة'): br[d]=('Z',9)
            elif d in ('القشرة','الدهانات'): br[d]=('B',ORDW[d])
            else: br[d]=('A',ORDW.get(d,9))
        for d in work:
            if p in RT and d in RT[p]:
                tot=sum(x[2] for x in RT[p][d]) or 1
                OPL[d]=[(n,m*q*(1-dsc)*(st['rem'][d]/tot)) for _,n,m in RT[p][d]]
    elif p in OVERRIDE:
        mn=ov_minutes(OVERRIDE[p],1); cur=ORDW.get(dept,0)
        work={d:m*q*(1-dsc) for d,m in mn.items() if ORDW.get(d,9)>=cur}
        if not work:
            work={dept: DEPT_AVG.get(dept,0.0)*q*(1-dsc)} if dept in DEPT_AVG else {}
            src=f'تقديري — متوسط قسم "{dept}"'
        for d,lst in OVERRIDE[p].items():
            if d in work: OPL[d]=[(o_[0],o_[2]*q*(1-dsc)) for o_ in lst]
        work=remap(work,p,cat_of(p))
        if 'نجارة النوم والسفرة' in work and 'نجارة التنجيد' in OPL:
            OPL['نجارة النوم والسفرة']=OPL.pop('نجارة التنجيد')
        br={d:('A',ORDW[d]) for d in work}
        src='مسار يدوي (منك)'
    elif p=='سبايدر كرسي سفرة':
        continue
    elif p in name2code:
        code=name2code[p]; _ol={}
        rm=remaining(code,dept,None if fs else stage,fs and dept not in ops[code],_ol)
        work={d:m*q*(1-dsc) for d,m in rm.items() if m>0}; br=branches(code)
        OPL={d:[(n,m*q*(1-dsc)) for n,m in l] for d,l in _ol.items() if d in work}
    else:
        k=cat_of(p)
        if not (k and k in catavg): continue
        _c=ORDW.get(dept,0)
        _base={d:m for d,m in catavg[k].items() if ORDW.get(d,9)>=_c}
        if not _base:
            # الأوردر واقف في قسم مش موجود في متوسط فئته — خده بمتوسط القسم نفسه
            _base={dept: DEPT_AVG.get(dept,0.0)} if dept in DEPT_AVG else dict(catavg[k])
            src=f'تقديري — متوسط قسم "{dept}" (فئة "{k}" مفيهاش القسم ده)'
        work=remap({d:m*q*(1-dsc) for d,m in _base.items()},p,k)
        br={d:('Z',9) if d.startswith('تشطيب') else (('B',ORDW[d]) if d in ('القشرة','الدهانات') else ('A',ORDW.get(d,9))) for d in work}
        src=f'تقديري — متوسط "{k}"'
    if not work: continue
    if 'اريكة' in p or 'اريكة' in trk: P,sub,lab,dd=1,0,'أريكة — تجاري عاجل',dt.date(2026,9,10)
    elif c and c in due:                P,sub,lab,dd=1,1,f'معرض — {cust.get(c,"")}',due[c]
    elif r[7]=='الاسبوع الحالي':        P,sub,lab,dd=1,2,'مصنع — الأسبوع الحالي',dt.date(2026,9,10)
    else:                               P,sub,lab,dd=2,3,'مصنع — باقي التشغيل',None
    work=bed_route(work,p)
    if is_bed(p):
        _o={}
        for d,l in OPL.items():
            d2=BED_DEPT if d in ('نجارة التنجيد','نجارة النوم والسفرة') else d
            _o[d2]=_o.get(d2,[])+l
        OPL=_o
        br={(BED_DEPT if d in ('نجارة التنجيد','نجارة النوم والسفرة') else d):v for d,v in br.items()}
    work={k2:v2/60.0 for k2,v2 in work.items()}      # دقيقة -> ساعة
    OPL={d:[(n,m/60.0) for n,m in l] for d,l in OPL.items()}
    ORDERS.append(dict(id=r[0],prod=p,qty=q,trk=trk,nf=c,P=P,sub=sub,lab=lab,due=dd,opl=OPL,work=dict(work),br=br,
                       src=(src+' | '+_note if _note else src),cur=dept,stage=stage,
                       ready=DEFERRED.get(p, MAT_DATE if fs else START)))
# سبايدر — أمر واحد مجمّع (128 كرسي، 90 خلصوا لحام)
r3=sum(m for n,m in chair['الاستانلس'][5:]); fl=sum(m for n,m in chair['الاستانلس'])
ORDERS.append(dict(id='SPIDER-128',prod='سبايدر كرسي سفرة (الكرسي الاستالس)',qty=128,trk='عمرو حسن',nf=None,
   P=1,sub=0,lab='عمرو حسن — 128 كرسي',due=dt.date(2026,9,10),
   work={'الاستانلس':(90*r3+38*fl)*.80/60,'السفنجة':128*sum(m for n,m in chair['السفنجة'])*.80/60,
         'التفصيل والكسوة':128*sum(m for n,m in chair['التفصيل والكسوة'])*.80/60},
   br={'الاستانلس':('C',1),'السفنجة':('A',2),'التفصيل والكسوة':('A',3)},
   src='ملف مسار منفصل',cur='الاستانلس',stage='تشطيب',ready=START,
   opl={'الاستانلس':[(n,(90 if i>=5 else 0)*m/60*.8+(38*m/60*.8)) for i,(n,m) in enumerate(chair['الاستانلس'])],
        'السفنجة':[(n,128*m/60*.8) for n,m in chair['السفنجة']],
        'التفصيل والكسوة':[(n,128*m/60*.8) for n,m in chair['التفصيل والكسوة']]}))
ORDERS.sort(key=lambda o:(o['P'], o['sub'], o['due'] or dt.date(2099,1,1)))
for i,o in enumerate(ORDERS,1): o['rank']=i
print(f"أوردرات داخلة الجدولة: {len(ORDERS)} | وحدات: {sum(o['qty'] for o in ORDERS)}")
print(f"  P1: {sum(1 for o in ORDERS if o['P']==1)} أمر / {sum(o['qty'] for o in ORDERS if o['P']==1)} وحدة")
print(f"  P2: {sum(1 for o in ORDERS if o['P']==2)} أمر / {sum(o['qty'] for o in ORDERS if o['P']==2)} وحدة")
print(f"  إجمالي ساعات: {sum(sum(o['work'].values()) for o in ORDERS):.0f}")

# ============ محرك الجدولة ============
# نسبة الشغل "المبكر" (تفصيل/خياطة) جوه التفصيل والكسوة — مش مستنية القطعة
def early_frac(o):
    p=o['prod']; code=name2code.get(p)
    if not code or 'التفصيل والكسوة' not in ops.get(code,{}): return 8/22
    l=ops[code]['التفصيل والكسوة']; tot=sum(x[2] for x in l)
    if tot<=0: return 8/22
    e=sum(x[2] for x in l if any(k in x[1] for k in ('تفصيل','خياط','تطريز')))
    return e/tot
for o in ORDERS: o['ef']=early_frac(o) if 'التفصيل والكسوة' in o['work'] else 0.0

def preds(o,d):
    """الأقسام اللي لازم تسبق القسم d لنفس الأوردر"""
    br=o['br']; 
    if d not in br: return []
    b,p=br[d]
    if b=='Z':   # التشطيب = تجميع كل الفروع
        last={}
        for dd,(bb,pp) in br.items():
            if bb=='Z' or dd not in o['work']: continue
            if bb not in last or pp>br[last[bb]][1]: last[bb]=dd
        return list(last.values())
    cand=[dd for dd,(bb,pp) in br.items() if bb==b and pp<p and dd in o['work']]
    return [max(cand,key=lambda x:br[x][1])] if cand else []

def run(use_ot=False, horizon=130):
    done=defaultdict(lambda: defaultdict(float))      # order -> dept -> ساعات اتعملت
    hist=defaultdict(lambda: defaultdict(list))       # order -> dept -> [(يوم, ساعات)]
    W0={o['rank']:o['work'] for o in ORDERS}
    deptday=defaultdict(lambda: defaultdict(float))   # يوم -> قسم -> ساعات مستخدمة
    plan=defaultdict(lambda: defaultdict(list))       # يوم -> قسم -> [(rank, ساعات)]
    for t,day in enumerate(DAYS[:horizon]):
        cap={d:WORKERS[d]*(NORM_H+(OT_H if use_ot else 0)) for d in WORKERS}
        sub={('التفصيل والكسوة','مبكر'):8*(NORM_H+(OT_H if use_ot else 0)),
             ('التفصيل والكسوة','متأخر'):14*(NORM_H+(OT_H if use_ot else 0))}
        for o in ORDERS:
            if day < o['ready']: continue
            k=o['rank']
            for d,Wd in o['work'].items():
                rem=Wd-done[k][d]
                if rem<=1e-6: continue
                # قيد السلف (تدفق قطعة قطعة)
                allow=rem
                ps=preds(o,d)
                if ps:
                    fr=1.0
                    for pd in ps:
                        wp_=o['work'].get(pd,0)
                        if wp_<=0: continue
                        cum=sum(h for dd,h in hist[k][pd] if dd<t)   # لحد نهاية إمبارح
                        lag=SEALER_LAG if pd=='الدهانات' else 0
                        if lag: cum=sum(h for dd,h in hist[k][pd] if dd<t-lag+0)
                        fr=min(fr,cum/wp_)
                    allow=min(allow, fr*Wd-done[k][d])
                if allow<=1e-6: continue
                if d=='التفصيل والكسوة':
                    We=Wd*o['ef']; Wl=Wd-We
                    de=min(max(0,We-min(done[k][d],We)), sub[(d,'مبكر')])
                    got=0
                    if de>1e-6: sub[(d,'مبكر')]-=de; got+=de
                    doneL=max(0,done[k][d]-We)
                    la=min(max(0,Wl-doneL), sub[(d,'متأخر')], max(0,allow-got))
                    if la>1e-6: sub[(d,'متأخر')]-=la; got+=la
                    alloc=min(got,rem)
                    cap[d]=max(0,cap[d]-alloc)
                else:
                    alloc=min(allow,rem,cap[d]); cap[d]-=alloc
                if alloc<=1e-6: continue
                done[k][d]+=alloc; hist[k][d].append((t,alloc))
                deptday[t][d]+=alloc; plan[t][d].append((k,alloc))
    return done,hist,deptday,plan

for label,ot in [('بالساعات العادية',False),('بالسهر الكامل',True)]:
    done,hist,deptday,plan=run(ot)
    fin={}
    for o in ORDERS:
        k=o['rank']; last=-1; ok=True
        for d,Wd in o['work'].items():
            if done[k][d] < Wd-1e-6: ok=False; break
            last=max(last,max(t for t,_ in hist[k][d]))
        fin[k]=DAYS[last] if ok and last>=0 else None
    n=sum(1 for v in fin.values() if v); tot=len(ORDERS)
    print(f"\n### {label}: خلّص {n}/{tot} أمر")
    if n:
        for P in (1,2):
            ds=[fin[o['rank']] for o in ORDERS if o['P']==P and fin[o['rank']]]
            if ds: print(f"   P{P}: آخر تسليم {max(ds)}  |  عدد اللي خلّص {len(ds)}")
    globals()['_last']=(done,hist,deptday,plan,fin)
