# -*- coding: utf-8 -*-
import openpyxl, re
from collections import defaultdict, Counter
U='/root/.claude/uploads/fdb1b8fa-74fd-56bd-9c54-782073111ccf/'
def num(v): return v if isinstance(v,(int,float)) else 0
def norm(s):
    s=str(s).strip().lower()
    s=re.sub('[إأآا]','ا',s); s=re.sub('[ىي]','ي',s); s=re.sub('ة','ه',s)
    s=re.sub(r'\(.*?\)','',s); s=re.sub(r'\s+',' ',s)
    return s.strip()

# الفئات — الترتيب مهم (الأخص الأول)
CATS=[('كرسي سفرة',['كرسي سفره','كراسي','كرسي']),
      ('كنبة ثلاثية',['كنبه ثلاثيه','ثلاثيه','كنب ثلاثيه']),
      ('كنبة ثنائية',['كنبه ثنائيه','ثنائيه']),
      ('ركنة',['ركنه','ركنة','حرف l','حرف u']),
      ('فوتيه',['فوتيه','فوتية','ارمشير']),
      ('بف',['بف','بانكت','بانكيت','رول']),
      ('سرير',['سرير']),
      ('كومود',['كومود']),
      ('تسريحة',['تسريحه','درسينج','دريسنج']),
      ('بوفيه',['بوفيه']),
      ('نيش',['نيش','فاترينه']),
      ('ترابيزة سفرة',['ترابيزه سفره','سفره']),
      ('ترابيزة جانبية',['ترابيزه جانبيه','سايد','ترابيزه صغيره','side','جانبيه']),
      ('ترابيزة رئيسية',['ترابيزه رئيسيه','ترابيزه كبيره','ديننج','dinning','ترابيزه','coffe','table']),
      ('كونصول',['كونصول','تي في','تى فى','تي فيونت','ديسك','desk']),
      ('مرايا',['مرايا','مرايه','برواز','mirror']),
      ('بلوك',['بلوك']),
      ]
def cat(name):
    n=norm(name)
    for c,keys in CATS:
        for k in keys:
            if k in n: return c
    return None

wr=openpyxl.load_workbook(U+'7b215713-_________________....xlsx',data_only=True)['Sheet1']
rrows=[r for r in wr.iter_rows(values_only=True)][1:]
prod=defaultdict(lambda: defaultdict(float)); pname={}
for r in rrows:
    if not r[3]: continue
    pname[r[3]]=str(r[5]).strip()
    if r[1]=='التغليف': continue
    prod[r[3]][r[1]]+=num(r[8])
name2code={}
for c,n in pname.items(): name2code.setdefault(n,c)

# متوسط كل فئة: متوسط دقائق كل قسم عبر منتجات الفئة اللي ليها مسار
catrows=defaultdict(list)
for code,depts in prod.items():
    c=cat(pname[code])
    if c: catrows[c].append(depts)
catavg={}
for c,lst in catrows.items():
    agg=defaultdict(float); n=len(lst)
    for d in lst:
        for dept,m in d.items(): agg[dept]+=m
    # القسم يتحسب لو موجود في نص المنتجات على الأقل
    present=Counter()
    for d in lst:
        for dept in d: present[dept]+=1
    catavg[c]={dept: agg[dept]/present[dept] for dept in agg if present[dept]>=max(1,n*0.5)}

wo=openpyxl.load_workbook(U+'70dac0e1-_____________________________.xlsx',data_only=True)['أوامر التصنيع']
orows=[r for r in list(wo.iter_rows(values_only=True))[1:] if any(v is not None for v in r)]
active=[r for r in orows if r[4] in ('IN_PROGRESS','NEW','PLANNED')]
missing=Counter()
for r in active:
    p=str(r[1]).strip() if r[1] else ''
    if p in ('Kosan','سبايدر كرسي سفرة') or p in name2code: continue
    missing[p]+=r[2] or 1

print("### متوسط كل فئة (ساعات/وحدة، بدون تغليف)")
for c in sorted(catavg,key=lambda x:-sum(catavg[x].values())):
    t=sum(catavg[c].values())
    print(f"  {c:16s} {t/60:6.2f} س   عيّنة {len(catrows[c]):3d} منتج   [{'، '.join(f'{d} {m/60:.1f}' for d,m in sorted(catavg[c].items(),key=lambda x:-x[1]))}]")

print("\n### المنتجات اللي مالهاش مسار → الفئة المقترحة")
ok=0; bad=[]
est_load=defaultdict(float)
for p,q in missing.most_common():
    c=cat(p)
    if c and c in catavg:
        ok+=q
        for d,m in catavg[c].items(): est_load[d]+=m*q
        if q>=4: print(f"  {q:4d}  {p:38s} → {c}  ({sum(catavg[c].values())/60:.2f} س/وحدة)")
    else: bad.append((p,q))
print(f"\n  اتغطّى بالتقدير: {ok} وحدة")
print(f"  لسه مش متصنّف : {sum(q for _,q in bad)} وحدة ({len(bad)} اسم)")
for p,q in sorted(bad,key=lambda x:-x[1])[:25]: print(f"     {q:4d}  {p!r}")
import json; json.dump({d:est_load[d] for d in est_load}, open('est_load.json','w'))
print("\n### الحِمل الإضافي من المنتجات المقدَّرة (ساعة)")
for d in sorted(est_load,key=lambda x:-est_load[x]): print(f"  {d:22s} {est_load[d]/60:7.0f} س")
