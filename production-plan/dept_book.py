# -*- coding: utf-8 -*-
"""تاب لكل قسم — على مستوى العملية: يبدأ من إيه، يقف عند إيه، يكمّل إمتى، يخلص إمتى"""
import runpy, math, datetime as dt, openpyxl
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
g=runpy.run_path('sched.py')
from status_update import STATUS
ORDERS=g['ORDERS']; DAYS=g['DAYS']; W=g['WORKERS']; NORM=g['NORM_H']; OT=g['OT_H']; run=g['run']
done,hist,deptday,plan=run(False)
doneO,histO,deptdayO,planO=run(True)
byrank={o['rank']:o for o in ORDERS}
AR=['الاتنين','الثلاثاء','الأربعاء','الخميس','الجمعة','السبت','الأحد']

def op_at(o,d,h):
    """عند إتمام h ساعة في القسم d — إحنا عند أنهي عملية وبأنهي نسبة"""
    l=o.get('opl',{}).get(d)
    if not l: return None,None
    c=0.0
    for i,(n,m) in enumerate(l):
        if h < c+m-1e-9:
            pct=(h-c)/m if m>0 else 0
            return n, pct
        c+=m
    return l[-1][0], 1.0
def label(o,d,h,is_end):
    n,pct=op_at(o,d,h)
    if n is None: return 'شغل القسم (مفيش تفصيل عمليات)'
    if is_end and pct>=0.999: return f'خلص «{n}» ✔'
    if pct<=0.001: return f'أول «{n}»'
    return f'«{n}» عند {pct*100:.0f}%'

# جدول: order -> dept -> [(t, ساعات)] مرتّب
seq=defaultdict(lambda: defaultdict(list))
for k in hist:
    for d,l in hist[k].items(): seq[k][d]=sorted(l)

wb=openpyxl.Workbook(); wb.remove(wb.active)
HF=Font(bold=True,color='FFFFFF',size=11); HB=PatternFill('solid',fgColor='1F4E79')
DAYF=Font(bold=True,size=12); DAYB=PatternFill('solid',fgColor='BDD7EE'); FRI=PatternFill('solid',fgColor='E2EFDA')
P1F=PatternFill('solid',fgColor='FFC7CE'); P2F=PatternFill('solid',fgColor='FFF2CC')
thin=Border(*[Side(style='thin',color='B0B0B0')]*4)

def tab(d,title):
    ws=wb.create_sheet(title); ws.sheet_view.rightToLeft=True
    ws.append([f'خطة شغل قسم: {d}',None,None,None,f'عدد العمال: {W[d]}',
               f'الطاقة: {W[d]*NORM:.0f} ساعة عادي + {W[d]*OT:.0f} سهر','','','','','',''])
    ws['A1'].font=Font(bold=True,size=14); ws.append([])
    ws.append(['التاريخ','اليوم','أولوية','كود الأمر','المنتج','كمية','العميل',
               'يبدأ اليوم من','يقف عند (آخر اليوم)','ساعات','عمال','يكمّل يوم','يخلص القسم يوم','ملاحظة'])
    for c in ws[3]: c.font=HF;c.fill=HB;c.alignment=Alignment(horizontal='center',wrap_text=True)
    row=4
    for t,day in enumerate(DAYS):
        items=plan[t][d]
        if not items or day>dt.date(2026,10,10): continue
        tot=sum(h for _,h in items); need=min(math.ceil(tot/NORM),W[d])
        exc=max(0.0,deptdayO[t][d]-W[d]*NORM); otp=min(W[d],math.ceil(exc/OT)) if exc>0.01 else 0
        ws.append([str(day),AR[day.weekday()],'— ملخص اليوم —','',f'إجمالي {tot:.1f} ساعة',
                   len(items),f'الطاقة {W[d]*NORM:.0f} س',f'التحميل {100*tot/(W[d]*NORM):.0f}%',
                   (f'{otp} يسهروا لتقديم الشغل' if otp else 'مش محتاج سهر'),
                   round(tot,1),need,'','',''])
        for c in ws[row]: c.font=DAYF; c.fill=FRI if day.weekday()==4 else DAYB; c.border=thin
        row+=1
        for k,h in sorted(items,key=lambda x:(byrank[x[0]]['P'],byrank[x[0]]['sub'],-x[1])):
            o=byrank[k]; L=seq[k][d]
            before=sum(hh for tt,hh in L if tt<t)
            after=before+h
            nxt=[tt for tt,_ in L if tt>t]
            last=DAYS[max(tt for tt,_ in L)]
            ws.append([str(day),AR[day.weekday()],f"P{o['P']}",o['id'],o['prod'],o['qty'],
                       (o['trk'] or '—')[:34],
                       label(o,d,before,False), label(o,d,after,True),
                       round(h,2), max(1,round(h/NORM,1)),
                       str(DAYS[nxt[0]]) if nxt else 'خلص ✔', str(last),
                       STATUS.get(o['prod'],{}).get('note','')])
            f=P1F if o['P']==1 else P2F
            for c in ws[row]: c.fill=f; c.border=thin
            row+=1
    for i,x in enumerate([12,10,7,17,28,6,34,30,30,9,8,13,14,40],1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=x
    ws.freeze_panes='A4'

TABS=[('نجارة التنجيد','1-نجارة التنجيد'),('التفصيل والكسوة','2-الكسوة'),('الدهانات','3-الدهانات'),
      ('نجارة النوم والسفرة','4-نجارة النوم والسفرة'),('السفنجة','5-السفنجة'),('القشرة','6-القشرة'),
      ('الاستانلس','7-الاستانلس'),('تشطيب التنجيد','8-تشطيب التنجيد'),('تشطيب نوم وسفرة','9-تشطيب نوم وسفرة')]
ws=wb.create_sheet('0-ملخص كل الأقسام'); ws.sheet_view.rightToLeft=True
ws.append(['ملخص يومي — كل قسم يشتغل كام ساعة وكام عامل']); ws['A1'].font=Font(bold=True,size=14); ws.append([])
ws.append(['التاريخ','اليوم']+[d for d,_ in TABS])
for c in ws[3]: c.font=HF;c.fill=HB;c.alignment=Alignment(horizontal='center',wrap_text=True)
r=4
for t,day in enumerate(DAYS):
    if not deptday[t] or day>dt.date(2026,10,10): continue
    line=[str(day),AR[day.weekday()]]
    for d,_ in TABS:
        h=deptday[t][d]
        line.append(f"{h:.0f} س / {min(math.ceil(h/NORM),W[d])} عامل" if h>0.05 else '—')
    ws.append(line)
    for c in ws[r]:
        c.border=thin
        if day.weekday()==4: c.fill=FRI
    r+=1
for i,x in enumerate([12,10]+[20]*len(TABS),1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=x
ws.freeze_panes='C4'
for d,title in TABS: tab(d,title)
wb.move_sheet('0-ملخص كل الأقسام', offset=-len(TABS))
wb.save('خطة_المشرفين.xlsx'); print("تم: خطة_المشرفين.xlsx")
