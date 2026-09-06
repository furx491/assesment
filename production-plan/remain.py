# -*- coding: utf-8 -*-
"""يحسب الشغل *الفاضل* لكل أوردر من مرحلته الحالية ورايح — مش المسار كامل."""
import openpyxl, re, difflib, json
from collections import defaultdict, Counter
U='/root/.claude/uploads/fdb1b8fa-74fd-56bd-9c54-782073111ccf/'
def num(v): return v if isinstance(v,(int,float)) else 0
def nz(s):
    s=str(s).strip().lower(); s=re.sub('[إأآا]','ا',s); s=re.sub('[ىي]','ي',s)
    s=re.sub('ة','ه',s); s=re.sub(r'\(.*?\)',' ',s); s=re.sub(r'[^\w\s]',' ',s)
    return re.sub(r'\s+',' ',s).strip()
def cat_of(n, CATS):
    x=nz(n)
    for c,ks in CATS:
        for k in ks:
            if k in x: return c
    return None
CATS=[('كرسي سفرة',['كرسي سفره','كراسي','كرسي']),('نيو بيكا كنبة',['بيكا كنبه مقاس','بيكا كنبه']),
      ('كنبة ثلاثية',['كنبه ثلاثيه','ثلاثيه']),('كنبة ثنائية',['كنبه ثنائيه','ثنائيه','double sofa']),
      ('ركنة',['ركنه','حرف l','حرف u']),('فوتيه',['فوتيه','فوتية','ارمشير']),
      ('بف',['بف','بانكت','بانكيت','رول']),('سرير',['سرير','bed']),('كومود',['كومود','commode']),
      ('تسريحة',['تسريحه','درسينج','دريسنج','درسيينج','نيش','جزامه','فاترينه']),
      ('بوفيه',['بوفيه','يوفيه']),('ترابيزة سفرة',['ترابيزه سفره','سفره']),
      ('ترابيزة جانبية',['ترابيزه جانبيه','سايد','ترابيزه صغيره','side','جانبيه']),
      ('ترابيزة رئيسية',['ترابيزه رئيسيه','ترابيزه كبيره','ديننج','dinning','ترابيزه','coffe','table']),
      ('كونصول',['كونصول','تي في','تى فى','تي فيونت','ديسك','desk']),
      ('مرايا',['مرايا','مرايه','برواز','mirror']),('بلوك',['بلوك']),('كنبة',['كنبه','sofa'])]

# ===== المسارات =====
wr=openpyxl.load_workbook(U+'7b215713-_________________....xlsx',data_only=True)['Sheet1']
ops=defaultdict(lambda: defaultdict(list)); pname={}
for r in [x for x in wr.iter_rows(values_only=True)][1:]:
    if not r[3]: continue
    pname[r[3]]=str(r[5]).strip()
    if r[1]=='التغليف': continue
    ops[r[3]][r[1]].append((r[6] if isinstance(r[6],int) else 999, str(r[7] or ''), num(r[8])))
for c in ops:
    for d in ops[c]: ops[c][d].sort()
name2code={}
for c,n in pname.items(): name2code.setdefault(n,c)

ORD={'نجارة التنجيد':1,'نجارة النوم والسفرة':1,'الفايبر جلاس':1,'الاستانلس':1,
     'السفنجة':2,'القشرة':2,'التفصيل والكسوة':3,'الدهانات':3,'تشطيب التنجيد':4,'تشطيب نوم وسفرة':4}
def branches(code):
    """يرجّع dict: قسم -> (فرع, ترتيب) حسب بنية المنتج"""
    d=set(ops[code]); res={}
    branched = ('السفنجة' in d and 'القشرة' in d)
    for dp in d:
        if dp in ('تشطيب التنجيد','تشطيب نوم وسفرة'): res[dp]=('Z',1)
        elif dp=='الاستانلس' and len(d)>1:          res[dp]=('C',1)
        elif branched and dp in ('القشرة','الدهانات'): res[dp]=('B',ORD[dp]-1)
        else: res[dp]=('A',ORD.get(dp,9))
    return res

def remaining(code, cur_dept, cur_stage, from_start):
    """دقائق فاضلة لكل قسم"""
    out=defaultdict(float); br=branches(code)
    if from_start or not cur_dept or cur_dept not in ops[code]:
        for dp,lst in ops[code].items(): out[dp]+=sum(x[2] for x in lst)
        return out,'من أول المسار'
    cb,cp=br[cur_dept]
    # القسم الحالي: من المرحلة الحالية ورايح
    lst=ops[code][cur_dept]; idx=0; how='القسم من أوله'
    if cur_stage:
        s=nz(cur_stage); best=None;bs=0
        for i,(sq,nm,mn) in enumerate(lst):
            n2=nz(nm)
            r = 1.0 if (s and (s in n2 or n2 in s)) else difflib.SequenceMatcher(None,s,n2).ratio()
            if r>bs: bs=r; best=i
        if bs>=0.60: idx=best; how=f'اتطابقت مع "{lst[best][1].strip()}" ({bs:.0%})'
    out[cur_dept]+=sum(x[2] for x in lst[idx:])
    for dp,lstd in ops[code].items():
        if dp==cur_dept: continue
        b,p=br[dp]
        if b=='Z' or b!=cb or (b==cb and p>cp): out[dp]+=sum(x[2] for x in lstd)
    return out,how

# ===== الأوردرات =====
DEPTNAME={'نجارة التنجيد','نجارة النوم والسفرة','التفصيل والكسوة','الدهانات','القشرة','السفنجة',
          'الاستانلس','الفايبر جلاس','دهانات','نجارة','التخطيط'}
wo=openpyxl.load_workbook(U+'70dac0e1-_____________________________.xlsx',data_only=True)['أوامر التصنيع']
orows=[r for r in list(wo.iter_rows(values_only=True))[1:] if any(v is not None for v in r)]
active=[r for r in orows if r[4] in ('IN_PROGRESS','NEW','PLANNED')]

# متوسطات الفئات (من المنتجات اللي ليها مسار)
catrows=defaultdict(list)
for c in ops:
    k=cat_of(pname[c],CATS)
    if k: catrows[k].append({d:sum(x[2] for x in l) for d,l in ops[c].items()})
catavg={};catn={}
for k,l in catrows.items():
    agg=defaultdict(float); pres=Counter()
    for d in l:
        for dp,m in d.items(): agg[dp]+=m; pres[dp]+=1
    catavg[k]={dp:agg[dp]/pres[dp] for dp in agg if pres[dp]>=max(1,len(l)*0.5)}; catn[k]=len(l)

def bdisc(q): return 0.20 if q>=20 else 0.15 if q>=10 else 0.10 if q>=5 else 0.0
qty=Counter()
for r in active:
    p=str(r[1]).strip() if r[1] else ''
    if p!='Kosan': qty[p]+=r[2] or 1

# مسار سبايدر
wc=openpyxl.load_workbook(U+'485a723f-______________________.xlsx',data_only=True)['Sheet1']
DM={'استالس':'الاستانلس','نجارة تنجيد':'نجارة التنجيد','سفنجة':'السفنجة','كسوة':'التفصيل والكسوة'}
chair=defaultdict(list)
for r in list(wc.iter_rows(values_only=True))[2:]:
    if r[0] and r[1] in DM: chair[DM[r[1]]].append((r[0].strip(),num(r[2])))

load=defaultdict(float); how_c=Counter(); unclass=Counter(); est_units=0
for r in active:
    p=str(r[1]).strip() if r[1] else ''; q=r[2] or 1
    if p=='Kosan': continue
    dsc=bdisc(qty[p])
    if p=='سبايدر كرسي سفرة':
        continue   # يتحسب مرة واحدة تحت
    stage=r[5]; dept=r[6]
    from_start = (r[4]=='NEW') or (stage is None) or (str(stage).strip() in DEPTNAME)
    if p in name2code:
        rem,how = remaining(name2code[p], dept, None if from_start else stage, from_start and dept not in ops[name2code[p]])
        how_c[how.split(' (')[0] if how.startswith('اتطابقت') else how]+=1
        for d,m in rem.items(): load[d]+=m*q*(1-dsc)
    else:
        k=cat_of(p,CATS)
        if k and k in catavg:
            est_units+=q
            for d,m in catavg[k].items(): load[d]+=m*q*(1-dsc)
        else: unclass[p]+=q
# سبايدر
rest3=sum(m for n,m in chair['الاستانلس'][5:]); full=sum(m for n,m in chair['الاستانلس'])
sp={'الاستانلس':90*rest3+38*full,'السفنجة':128*sum(m for n,m in chair['السفنجة']),
    'التفصيل والكسوة':128*sum(m for n,m in chair['التفصيل والكسوة'])}
for d,m in sp.items(): load[d]+=m*0.80

EFF=.75; BUF=.10
norm_h=(5*1.00+3*0.90)*EFF*(1-BUF); ot_h=4.5*0.70*EFF*(1-BUF)
W={'نجارة التنجيد':42,'السفنجة':19,'التفصيل والكسوة':22,'الاستانلس':16,'الدهانات':30,
   'القشرة':9,'نجارة النوم والسفرة':15,'تشطيب التنجيد':5,'تشطيب نوم وسفرة':3}
print(f"{'القسم':<22}{'عمال':>5}{'طاقة/يوم':>10}{'الفاضل(س)':>11}{'أيام':>7}{'أسابيع':>8}")
res=[]
for d in sorted(W,key=lambda x:-load[x]/(W[x]*norm_h)):
    cap=W[d]*norm_h; days=load[d]/60/cap
    print(f"{d:<22}{W[d]:>5}{cap:>10.1f}{load[d]/60:>11.0f}{days:>7.1f}{days/6:>8.1f}")
    res.append([d,W[d],round(cap,1),round(load[d]/60),round(days,1),round(days/6,1),round(W[d]*ot_h,1)])
print("\n### إزاي حدّدت نقطة البداية")
for k,v in how_c.most_common(8): print(f"  {v:5d}  {k}")
print(f"\nوحدات بوقت تقديري: {est_units} | لسه مالهاش تقدير: {sum(unclass.values())} ({len(unclass)} اسم)")
for p,q in unclass.most_common(): print(f"   {q:3d}  {p!r}")
json.dump({'res':res,'load':dict(load),'unclass':dict(unclass),'catavg':{k:dict(v) for k,v in catavg.items()},'catn':catn},open('remain.json','w'))
