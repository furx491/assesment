# -*- coding: utf-8 -*-
import openpyxl, re, json
from override import OVERRIDE, ov_minutes
# خلصوا فعلاً — بأمر المستخدمة (الترابيزات بس، الركن لأ)
EXCLUDE={'Kosan','تانجل ترابيزه جانبيه','هاربر ترابيزة رئيسية اريكة','لوليتا ترابيزة جانبية'}
from collections import defaultdict, Counter
exec(open('remain.py',encoding='utf-8').read().split('# ===== الأوردرات =====')[0])

U='/root/.claude/uploads/fdb1b8fa-74fd-56bd-9c54-782073111ccf/'
DEPTNAME={'نجارة التنجيد','نجارة النوم والسفرة','التفصيل والكسوة','الدهانات','القشرة','السفنجة',
          'الاستانلس','الفايبر جلاس','دهانات','نجارة','التخطيط'}
wo=openpyxl.load_workbook(U+'70dac0e1-_____________________________.xlsx',data_only=True)['أوامر التصنيع']
orows=[r for r in list(wo.iter_rows(values_only=True))[1:] if any(v is not None for v in r)]
active=[r for r in orows if r[4] in ('IN_PROGRESS','NEW','PLANNED')]
# خطة المعرض
wp=openpyxl.load_workbook(U+'738ab119-___________________.xlsx',data_only=True)['خطة الانتاج ']
DAYS={'السبت':'2026-09-05','الاحد':'2026-09-06','الأحد':'2026-09-06','الاتنين':'2026-09-07',
      'الثلاثاء':'2026-09-08','الأربعاء':'2026-09-09','الاربعاء':'2026-09-09','الخميس':'2026-09-10'}
duedate={}; cust={}; note={}
day=None
for r in wp.iter_rows(min_row=5,values_only=True):
    d=str(r[1] or '').split('\n')[0].strip()
    if d in DAYS: day=DAYS[d]
    for m in re.findall(r'NF-?(\d{3,6})', str(r[2] or '')):
        c='NF-'+m.zfill(6); duedate.setdefault(c,day); cust.setdefault(c,str(r[3] or '').strip()); note.setdefault(c,str(r[6] or '').strip())
def nfcode(t):
    m=re.search(r'[Nn][Ff]-?(\d{3,6})',str(t or ''));  return 'NF-'+m.group(1).zfill(6) if m else None

catrows=defaultdict(list)
for c in ops:
    k=cat_of(pname[c],CATS)
    if k: catrows[k].append({d:sum(x[2] for x in l) for d,l in ops[c].items()})
catavg={};catn={}
for k,l in catrows.items():
    agg=defaultdict(float); pres=Counter()
    for d in l:
        for dp,m in d.items(): agg[dp]+=m; pres[dp]+=1
    catavg[k]={dp:agg[dp]/pres[dp] for dp in agg if pres[dp]>=max(1,len(l)*.5)}; catn[k]=len(l)
def bdisc(q): return .20 if q>=20 else .15 if q>=10 else .10 if q>=5 else 0.0
qty=Counter()
for r in active:
    p=str(r[1]).strip() if r[1] else ''
    if p not in EXCLUDE: qty[p]+=r[2] or 1

def prio(r):
    p=str(r[1] or ''); t=str(r[3] or ''); c=nfcode(t)
    if 'اريكة' in p or 'اريكة' in t: return ('P1','أريكة — تجاري عاجل','2026-09-10')
    if p=='سبايدر كرسي سفرة':        return ('P1','عمرو حسن — 128 كرسي','2026-09-10')
    if c and c in duedate:            return ('P2',f'معرض — {cust.get(c,"")} ({c})',duedate[c])
    if r[7]=='الاسبوع الحالي':        return ('P2','مصنع — الأسبوع الحالي','2026-09-10')
    return ('P3','مصنع — بدون تاريخ','')

ORDW={'نجارة التنجيد':1,'نجارة النوم والسفرة':1,'الاستانلس':1,'السفنجة':2,'القشرة':2,
      'التفصيل والكسوة':3,'الدهانات':3,'تشطيب التنجيد':4,'تشطيب نوم وسفرة':4}
def remwork(r):
    p=str(r[1]).strip() if r[1] else ''; q=r[2] or 1; d=bdisc(qty[p]); out=defaultdict(float)
    if p=='سبايدر كرسي سفرة': return out
    if p in OVERRIDE:
        mn=ov_minutes(OVERRIDE[p],1); cur=ORDW.get(r[6],0)
        for dp,m in mn.items():
            if ORDW.get(dp,9)>=cur: out[dp]+=m*q*(1-d)     # من القسم الحالي ورايح
        return out
    if p in name2code:
        code=name2code[p]; stage=r[5]; dept=r[6]
        fs=(r[4]=='NEW') or (stage is None) or (str(stage).strip() in DEPTNAME)
        rem,_=remaining(code,dept,None if fs else stage, fs and dept not in ops[code])
        for k,m in rem.items(): out[k]+=m*q*(1-d)
    else:
        k=cat_of(p,CATS)
        if k and k in catavg:
            for dp,m in catavg[k].items(): out[dp]+=m*q*(1-d)
    return out

wc=openpyxl.load_workbook(U+'485a723f-______________________.xlsx',data_only=True)['Sheet1']
DM={'استالس':'الاستانلس','نجارة تنجيد':'نجارة التنجيد','سفنجة':'السفنجة','كسوة':'التفصيل والكسوة'}
chair=defaultdict(list)
for r in list(wc.iter_rows(values_only=True))[2:]:
    if r[0] and r[1] in DM: chair[DM[r[1]]].append((r[0].strip(),num(r[2])))
r3=sum(m for n,m in chair['الاستانلس'][5:]); fl=sum(m for n,m in chair['الاستانلس'])

grp=defaultdict(lambda: defaultdict(float)); units=Counter(); glabel={}
for r in active:
    if str(r[1] or '').strip() in EXCLUDE: continue
    P,lab,due=prio(r); key=(P,due,lab); units[key]+=r[2] or 1; glabel[key]=lab
    for d,m in remwork(r).items(): grp[key][d]+=m
kspider=('P1','2026-09-10','عمرو حسن — 128 كرسي')
for d,m in {'الاستانلس':90*r3+38*fl,'السفنجة':128*sum(m for n,m in chair['السفنجة']),
            'التفصيل والكسوة':128*sum(m for n,m in chair['التفصيل والكسوة'])}.items():
    grp[kspider][d]+=m*0.80

EFF=.75;BUF=.10; norm_h=(5*1+3*.9)*EFF*(1-BUF)
W={'نجارة التنجيد':42,'السفنجة':19,'التفصيل والكسوة':22,'الاستانلس':16,'الدهانات':30,
   'القشرة':9,'نجارة النوم والسفرة':15,'تشطيب التنجيد':5,'تشطيب نوم وسفرة':3}
rows=sorted(grp,key=lambda k:(k[0],k[1] or '9999'))
print(f"{'أولوية':<4}{'التاريخ':<12}{'المجموعة':<38}{'وحدات':>6}{'ساعات':>8}{'أبطأ قسم':>22}{'أيام':>6}")
out=[]
for k in rows:
    tot=sum(grp[k].values())/60
    worst=max(grp[k],key=lambda d: grp[k][d]/60/(W[d]*norm_h)) if grp[k] else '—'
    wd=grp[k][worst]/60/(W[worst]*norm_h) if grp[k] else 0
    print(f"{k[0]:<4}{(k[1] or '—'):<12}{glabel[k][:36]:<38}{units[k]:>6}{tot:>8.0f}{worst:>22}{wd:>6.1f}")
    out.append([k[0],k[1] or '—',glabel[k],units[k],round(tot),worst,round(wd,1),
                ' | '.join(f'{d} {grp[k][d]/60:.0f}س' for d in sorted(grp[k],key=lambda x:-grp[k][x]))])
print()
for P in ['P1','P2','P3']:
    ks=[k for k in rows if k[0]==P]; t=sum(sum(grp[k].values()) for k in ks)/60; u=sum(units[k] for k in ks)
    print(f"  {P}: {u:4d} وحدة | {t:6.0f} ساعة")
json.dump(out,open('prio.json','w'))
